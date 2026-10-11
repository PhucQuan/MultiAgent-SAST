from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


TABLE_SEPARATOR_RE = re.compile(r"^\|(?:\s*:?-{3,}:?\s*\|)+\s*$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
ORDERED_RE = re.compile(r"^\d+\.\s+(.*)$")
INLINE_TOKEN_RE = re.compile(r"(\*\*.*?\*\*|`.*?`)")


from docx.shared import Pt, RGBColor


BLACK = RGBColor(0, 0, 0)
FONT_PRIMARY = "Times New Roman"
FONT_CODE = "Consolas"


def clear_document(document: Document) -> None:
    body = document._body._element
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def normalize_inline(text: str) -> str:
    return text.replace("\\|", "|").strip()


def add_inline_runs(paragraph: Paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE_TOKEN_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            run.font.color.rgb = BLACK
            run.font.name = FONT_PRIMARY
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.color.rgb = BLACK
            run.font.name = FONT_PRIMARY
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = FONT_CODE
            run.font.color.rgb = BLACK
        else:
            run = paragraph.add_run(token)
            run.font.color.rgb = BLACK
            run.font.name = FONT_PRIMARY
        cursor = match.end()

    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        run.font.color.rgb = BLACK
        run.font.name = FONT_PRIMARY


def split_table_row(line: str) -> list[str]:
    stripped = line.strip().strip("|")
    return [normalize_inline(cell) for cell in stripped.split("|")]


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return

    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"

    for row_idx, row in enumerate(rows):
        for col_idx in range(column_count):
            value = row[col_idx] if col_idx < len(row) else ""
            cell = table.cell(row_idx, col_idx)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            add_inline_runs(paragraph, value)


def add_heading(document: Document, level: int, text: str, title_written: bool) -> bool:
    clean_text = normalize_inline(text)
    if level == 1 and not title_written:
        paragraph = document.add_paragraph(style="Title")
        add_inline_runs(paragraph, clean_text)
        return True

    heading_level = min(max(level - 1, 1), 9)
    paragraph = document.add_paragraph(style=f"Heading {heading_level}")
    add_inline_runs(paragraph, clean_text)
    return title_written


def add_paragraph(document: Document, text: str, style: str = "Normal") -> None:
    paragraph = document.add_paragraph(style=style)
    add_inline_runs(paragraph, normalize_inline(text))


def is_table_start(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines):
        return False
    return lines[index].lstrip().startswith("|") and TABLE_SEPARATOR_RE.match(lines[index + 1].strip()) is not None


def parse_markdown(document: Document, markdown_text: str) -> None:
    lines = markdown_text.splitlines()
    title_written = False
    index = 0

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if is_table_start(lines, index):
            rows: list[list[str]] = [split_table_row(lines[index])]
            index += 2
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                rows.append(split_table_row(lines[index]))
                index += 1
            add_table(document, rows)
            continue

        heading_match = HEADING_RE.match(stripped)
        if heading_match:
            level = len(heading_match.group(1))
            title_written = add_heading(document, level, heading_match.group(2), title_written)
            index += 1
            continue

        if stripped.startswith("- "):
            while index < len(lines) and lines[index].strip().startswith("- "):
                add_paragraph(document, lines[index].strip()[2:], style="List Bullet")
                index += 1
            continue

        ordered_match = ORDERED_RE.match(stripped)
        if ordered_match:
            while index < len(lines):
                current = lines[index].strip()
                current_match = ORDERED_RE.match(current)
                if not current_match:
                    break
                add_paragraph(document, current_match.group(1), style="List Number")
                index += 1
            continue

        if stripped.startswith("**") and stripped.endswith("**") and stripped.count("**") == 2:
            add_paragraph(document, stripped)
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate:
                break
            if HEADING_RE.match(candidate) or candidate.startswith("- ") or ORDERED_RE.match(candidate) or is_table_start(lines, index):
                break
            paragraph_lines.append(candidate)
            index += 1

        add_paragraph(document, " ".join(paragraph_lines))


def enforce_black_theme(document: Document) -> None:
    # Update style definitions
    for style_name in [
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Heading 4",
        "Heading 5",
        "Heading 6",
        "Normal",
        "List Bullet",
        "List Number",
    ]:
        if style_name in document.styles:
            style = document.styles[style_name]
            try:
                style.font.color.rgb = BLACK
                style.font.name = FONT_PRIMARY
            except Exception:
                pass

    # Ensure every single run in paragraphs is black
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = BLACK
            if run.font.name != FONT_CODE:
                run.font.name = FONT_PRIMARY

    # Ensure every single run in tables is black
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = BLACK
                        if run.font.name != FONT_CODE:
                            run.font.name = FONT_PRIMARY


def export_markdown_to_docx(markdown_path: Path, output_path: Path, template_path: Path | None = None) -> None:
    if template_path and template_path.exists():
        document = Document(str(template_path))
    elif output_path.exists():
        document = Document(str(output_path))
    else:
        document = Document()

    clear_document(document)
    parse_markdown(document, markdown_path.read_text(encoding="utf-8"))
    enforce_black_theme(document)
    document.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export Markdown thesis docs to DOCX.")
    parser.add_argument("markdown_files", nargs="+", help="Markdown files to export.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    for markdown_file in args.markdown_files:
        markdown_path = Path(markdown_file).resolve()
        output_path = markdown_path.with_suffix(".docx")
        export_markdown_to_docx(markdown_path, output_path, template_path=output_path if output_path.exists() else None)
        print(f"Exported {markdown_path.name} -> {output_path.name}")


if __name__ == "__main__":
    main()
